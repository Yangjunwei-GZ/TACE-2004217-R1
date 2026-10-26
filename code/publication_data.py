"""Create tables and reporting source data from completed validation, without refitting."""
from pathlib import Path
import os,json,gzip,sys
sys.path.insert(0,str(Path(__file__).resolve().parent))
from validation_integrity import audit_validation, verify_paired_summary
import numpy as np
import pandas as pd
from lifelines import KaplanMeierFitter
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
ROOT=Path(os.environ.get('TACE_PROJECT_ROOT','.'))/'revision_R1_20261014'; OUT=ROOT/'05_figures_and_tables'

LABELS={
 'clinical':'Clinical', 'radiomics_only':'Radiomics only',
 'clinical_diameter':'Clinical + maximum 3D diameter',
 'clinical_selected_volume':'Clinical + selected physical volume',
 'clinical_total_volume':'Clinical + total segmented physical volume',
 'clinical_mean_hu':'Clinical + mean CT intensity', 'clinical_median_hu':'Clinical + median CT intensity',
 'clinical_p10_hu':'Clinical + 10th percentile CT intensity',
 'clinical_diameter_mean_hu':'Clinical + diameter + mean CT intensity',
 'clinical_diameter_p10_hu':'Clinical + diameter + 10th percentile CT intensity',
 'clinical_axial_diameter':'Clinical + clinical axial diameter', 'clinical_legacy_seven':'Historical seven-variable clinical baseline',
 'clinical_age_spline':'Clinical with three-df age spline',
 'clinical_size_intensity_RSF':'Clinical + diameter + mean intensity (survival forest)',
 'clinical_size_intensity_radiomics_RSF':'Clinical + diameter + mean intensity + radiomics (survival forest)'
}
for k,v in list(LABELS.items()):
 if k in ['clinical','clinical_diameter','clinical_selected_volume','clinical_total_volume','clinical_diameter_mean_hu','clinical_diameter_p10_hu']: LABELS[k+'_radiomics']=v+' + radiomics'

def value(v):
    if isinstance(v,(bool,np.bool_)): return 'Yes' if v else 'No'
    if isinstance(v,(list,dict)): return json.dumps(v,ensure_ascii=False)
    if pd.isna(v): return 'Not estimable' if isinstance(v,float) else ''
    return v.item() if isinstance(v,np.generic) else v

def add_sheet(wb,name,title,df,note=None):
    s=wb.create_sheet(name); s.append([title]); s.merge_cells(start_row=1,start_column=1,end_row=1,end_column=max(1,len(df.columns))); s.cell(1,1).font=Font(name='Times New Roman',size=12,bold=True)
    s.append(list(df.columns))
    for record in df.itertuples(index=False,name=None): s.append([value(v) for v in record])
    gray=Side(style='thin',color='D9D9D9')
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='Times New Roman',size=10,bold=c.row==2,color='FFFFFF' if c.row==2 else '000000')
            c.fill=PatternFill('solid',fgColor='455A64' if c.row==2 else 'FFFFFF' if c.row%2 else 'F2F4F5')
            c.alignment=Alignment(vertical='center',wrap_text=True); c.border=Border(left=gray,right=gray,top=gray,bottom=gray)
            if isinstance(c.value,float): c.number_format='0.000'
    for i,col in enumerate(df.columns,1): s.column_dimensions[get_column_letter(i)].width=min(48,max(16,len(str(col))+2))
    for row in s.iter_rows(min_row=2):
        lines=max(max(1,int(np.ceil(len(str(c.value or ''))/s.column_dimensions[c.column_letter].width))) for c in row)
        s.row_dimensions[row[0].row].height=min(220,max(24,lines*12+8))
    s.freeze_panes='A3'; s.auto_filter.ref=f'A2:{get_column_letter(len(df.columns))}{len(df)+2}'
    if note: s.append([]); s.append([note]); s.merge_cells(start_row=s.max_row,start_column=1,end_row=s.max_row,end_column=max(1,len(df.columns)))
    return s

def main():
    validation=audit_validation(ROOT)
    verify_paired_summary(ROOT,validation)
    OUT.mkdir(exist_ok=True)
    data=pd.read_csv(ROOT/'03_model_development/FINAL_REVISION_PATIENT_SCORES.csv',dtype={'patient_id':str}); perf=pd.read_csv(ROOT/'04_validation/full_pipeline_performance.csv'); delta=pd.read_csv(ROOT/'04_validation/paired_incremental_performance.csv')
    perf['Model']=perf.model.map(LABELS); delta['Baseline']=delta.baseline.map(LABELS)
    perf.to_csv(OUT/'SOURCE_MODEL_PERFORMANCE.csv',index=False); delta.to_csv(OUT/'SOURCE_PAIRED_INCREMENTAL_VALUE.csv',index=False)
    records=[]; predictions=[]; ibs=[]; tuning=[]
    for p in sorted((ROOT/'04_validation/checkpoints').glob('point_*.json.gz')):
        with gzip.open(p,'rt',encoding='utf-8') as h:r=json.load(h)
        for f in r['folds']:
            dev=f.get('radiomics_development',{})
            for x in dev.get('tuning',[]): tuning.append(dict(run=r['run_id'],fold=f['fold'],**x))
            for model,s in f['models'].items():
                for m in s.get('time_metrics',[]):
                    if m['time']=='IBS':
                        if m['status']=='evaluated': ibs.append(dict(run=r['run_id'],fold=f['fold'],model=model,IBS=m['IBS'],grid_start=m['grid'][0],grid_end=m['grid'][-1],grid_points=len(m['grid']),n=m['n']))
                        continue
                    rec=dict(run=r['run_id'],fold=f['fold'],model=model,**{k:v for k,v in m.items() if k!='predicted_risk'}); records.append(rec)
                    if m['status']=='evaluated':
                        for pid,risk in zip(f['validation_ids'],m['predicted_risk']): predictions.append(dict(patient_id=pid,run=r['run_id'],fold=f['fold'],model=model,time=m['time'],predicted_risk=risk))
    detail=pd.DataFrame(records); detail.to_csv(OUT/'cross_validated_time_metric_details.csv',index=False)
    time_summary=detail[detail.status=='evaluated'].groupby(['model','time']).agg(AUC=('AUC','mean'),Brier=('Brier','mean'),evaluated_folds=('AUC','count'),total_validation_evaluations=('n','sum')).reset_index()
    time_summary.to_csv(OUT/'cross_validated_time_performance.csv',index=False); pd.DataFrame(ibs).to_csv(OUT/'cross_validated_monthly_IBS.csv',index=False)
    ibs_frame=pd.DataFrame(ibs)
    ibs_summary=ibs_frame.groupby('model').agg(IBS_mean=('IBS','mean'),IBS_median=('IBS','median'),IBS_SD=('IBS','std'),evaluated_folds=('IBS','count'),total_validation_evaluations=('n','sum'),grid_start_min_month=('grid_start','min'),grid_start_max_month=('grid_start','max'),grid_end_min_month=('grid_end','min'),grid_end_max_month=('grid_end','max'),grid_points_min=('grid_points','min'),grid_points_max=('grid_points','max')).reset_index()
    ibs_summary['interpretation']='Descriptive mean across fold-specific common-support monthly intervals; not a fixed 1-36 month IBS or CI'
    ibs_summary.to_csv(OUT/'cross_validated_monthly_IBS_summary.csv',index=False)
    pred=pd.DataFrame(predictions); pred.to_csv(OUT/'cross_validated_fixed_time_predictions.csv',index=False)
    avg=pred.groupby(['patient_id','model','time']).predicted_risk.mean().reset_index(); avg=avg.merge(data[['patient_id','os_months','death_event']],on='patient_id',validate='many_to_one')
    calibration=[]
    for (model,t),g in avg.groupby(['model','time']):
        g=g.copy(); g['bin']=pd.qcut(g.predicted_risk,5,labels=False,duplicates='drop')+1
        for b,group in g.groupby('bin'):
            km=KaplanMeierFitter().fit(group.os_months,group.death_event); observed=1-float(km.predict(t)); mean=group.predicted_risk.mean()
            calibration.append(dict(model=model,time_months=t,bin=int(b),n=len(group),mean_predicted_risk=mean,median_predicted_risk=group.predicted_risk.median(),observed_km_risk=observed,absolute_calibration_error=abs(mean-observed),method='Patient-averaged cross-validated risk; five quantile groups; observed Kaplan-Meier event risk'))
    cal=pd.DataFrame(calibration); cal.to_csv(OUT/'cross_validated_calibration.csv',index=False)
    cal_summary=cal.groupby(['model','time_months']).apply(lambda g:pd.Series({'MAE_unweighted':g.absolute_calibration_error.mean(),'MAE_n_weighted':np.average(g.absolute_calibration_error,weights=g.n),'patients':g.n.sum()}),include_groups=False).reset_index(); cal_summary.to_csv(OUT/'cross_validated_calibration_summary.csv',index=False)
    raw=pd.read_csv(ROOT/'01_data_audit/source_clinical_233.csv',dtype={'patient_id':str}); included=raw.patient_id.isin(data.patient_id)
    descriptors=[]
    for variable in ['age','gender_woman','cps','bclc','lab_afp','lesions_number','lesion1_diameter','lab_albumin','lab_bilirubin','tace_number','survival_time']:
        x=pd.to_numeric(raw.loc[included,variable],errors='coerce'); y=pd.to_numeric(raw.loc[~included,variable],errors='coerce'); denominator=np.sqrt((x.var()+y.var())/2)
        descriptors.append(dict(variable=variable,included_n=x.notna().sum(),excluded_n=y.notna().sum(),included_mean=x.mean(),excluded_mean=y.mean(),included_median=x.median(),excluded_median=y.median(),SMD=(x.mean()-y.mean())/denominator if denominator>0 else np.nan,included_missing=x.isna().sum(),excluded_missing=y.isna().sum(),notes='Total TACE sessions are post-baseline; albumin units in original definitions conflict with observed scale' if variable in ['tace_number','lab_albumin'] else 'Descriptive source-cohort comparison'))
    compare=pd.DataFrame(descriptors); compare.to_csv(OUT/'included_excluded_comparison.csv',index=False)
    f=pd.read_csv(ROOT/'03_model_development/FINAL_REVISION_MODEL_FREEZE.csv',dtype={'beta':str,'mean':str,'SD_population':str,'median_imputation':str})
    summary=json.loads((ROOT/'03_model_development/final_analysis_summary.json').read_text()); a=pd.read_csv(ROOT/'01_data_audit/phase_pairing_audit.csv'); eligibility=pd.read_csv(ROOT/'01_data_audit/phase_analysis_eligibility.csv')
    items=[{'Item':k,'Value':v} for k,v in dict(source_patients=len(raw),source_masks=len(a),primary_phase='Arterial',analysis_patients=summary['n'],deaths=summary['events'],censored=summary['censored'],analysis_lesions=summary['lesions'],original_features=107,KM_median_OS_months=summary['median_KM_OS'],reverse_KM_followup_months=summary['median_reverse_KM_followup']).items()]
    for col,label in [('age','Age (years)'),('lab_afp','AFP (ng/mL)'),('lesions_number','Source clinical lesion number'),('selected_diameter_mm','Selected maximum 3D diameter (mm)'),('selected_volume_cm3','Selected physical volume (cm3)'),('total_segmented_volume_cm3','Total segmented union volume (cm3)')]:
        x=data[col];items.append({'Item':label+', median (IQR)','Value':f'{x.median():.2f} ({x.quantile(.25):.2f}-{x.quantile(.75):.2f})'})
    for col,label in [('sex_female','Female sex'),('child_b','Child-Pugh B'),('bclc_b','BCLC B')]:
        items.append({'Item':label+', n (%)','Value':f'{int(data[col].sum())} ({data[col].mean()*100:.1f})'})
    cohort=pd.DataFrame(items)
    cohort['Item']=cohort.Item.replace({'source_patients':'Source patients','source_masks':'Source tumour masks','primary_phase':'CT contrast phase','analysis_patients':'Analysis patients','deaths':'Deaths','censored':'Censored','analysis_lesions':'Eligible arterial lesion pairs','original_features':'Original radiomics features','KM_median_OS_months':'Kaplan-Meier median OS (months)','reverse_KM_followup_months':'Reverse Kaplan-Meier median follow-up (months)'})
    wb=Workbook(); wb.remove(wb.active)
    add_sheet(wb,'Table 1','Table 1. Cohort and outcome characteristics.',cohort)
    specification=f"{summary['final_kind']}, penalty {summary['final_penalty']:g}"
    add_sheet(wb,'Table 2','Table 2. Final revision radiomics score reconstruction.',f[['feature','beta','mean','SD_population']],f'Population SD (ddof=0); coefficients are conditioned on the final full-cohort fit. Full-development performance comes from training-only validation. Actual selected specification: {specification}; full reconstruction metadata in Supplementary Table S12.')
    core=perf[perf.model.isin(['clinical','radiomics_only','clinical_diameter','clinical_diameter_mean_hu','clinical_diameter_mean_hu_radiomics'])].copy()
    core['row_order']=core.model.map({k:i for i,k in enumerate(['clinical','radiomics_only','clinical_diameter','clinical_diameter_mean_hu','clinical_diameter_mean_hu_radiomics'])});core=core.sort_values('row_order')
    add_sheet(wb,'Table 3','Table 3. Full-pipeline internally validated model performance.',core[['Model','cross_validated_C_index','CI_lower','CI_upper','repeat_mean_C_index','bootstrap_complete']],f"Primary C-index: predefined run 0, pair-weighted within-fold grouped cross-validation. Intervals: {validation['counts']['bootstrap']['planned']} patient bootstrap repetitions of the same randomized full CV estimator. Ten-repeat mean is split sensitivity, and has no attached interval. No optimism-corrected label is assigned. All {len(validation['plan']['model_names'])} applicable specifications are reported in Supplementary Tables S13-S14.")
    add_sheet(wb,'Table 4','Table 4. Paired radiomics increments beyond simple measurement baselines.',delta[['Baseline','paired_delta_C_index','CI_lower','CI_upper','bootstrap_joint_complete','ci_crosses_zero']],'Same patients, resamples, folds and evaluation sets within each contrast. Primary comparison: clinical + maximum 3D diameter + mean CT intensity, with versus without developed radiomics.')
    wb.save(OUT/'Main_Tables.xlsx')
    sw=Workbook(); sw.remove(sw.active)
    ph=pd.read_csv(ROOT/'03_model_development/PH_DIAGNOSTICS.csv'); ph['software']='lifelines; conditional penalised diagnostic fit'
    phglobal=pd.read_csv(ROOT/'03_model_development/PH_GLOBAL_RISK_SET_DIAGNOSTICS.csv').rename(columns={'chisq':'test_statistic','p':'P'}); phglobal['software']='R survival; conditional unpenalised Cox; global and covariate tests'
    ph=pd.concat([ph,phglobal],ignore_index=True)
    tables=[
      ('S1','Source mask and CT/phase pairing audit',a[['sample_id','patient_id','ct_phase','metadata_match','eligible','exclusion_reason','physical_volume_mm3']]),
      ('S2','Original clinical covariate availability',pd.read_csv(ROOT/'01_data_audit/clinical_variable_audit.csv')),
      ('S3','Historical versus revised validation definitions',pd.read_csv(ROOT/'01_data_audit/legacy_pipeline_audit.csv')),
      ('S4','Size/intensity collinearity',pd.read_csv(ROOT/'03_model_development/SIZE_INTENSITY_COLLINEARITY.csv')),
      ('S5','Training-only candidate tuning',pd.DataFrame(tuning)),
      ('S6','Nested cross-validation bootstrap summary records',pd.read_csv(ROOT/'04_validation/validation_run_metrics.csv')),
      ('S7','Cross-validated IPCW AUC and Brier scores',time_summary),
      ('S8','Cross-validated internal grouped calibration',cal),
      ('S9','Physical-volume selection audit; historical DCA retired',data[['patient_id','selected_volume_cm3','total_segmented_volume_cm3','n_segmented_arterial_lesions','lesions_number','lesion_selection_changed_from_voxels']]),
      ('S10','Included versus excluded source patients',compare),
      ('S11','Physical geometry and eligibility',a[['sample_id','patient_id','geometry_exact','geometry_tiny_tolerance','roi_in_image_physical_bounds','geometry_max_spacing_difference','geometry_max_origin_difference','exclusion_reason']]),
      ('S12','Final-model reconstruction and extraction settings',f),
      ('S13','Full-development performance intervals',perf),
      ('S14','Fair tumour size, volume and simple intensity model comparisons',perf),
      ('S15','Patient-level source-phase eligibility',eligibility),
      ('S16','Cox assumption diagnostics',ph),
      ('S17','Paired complete-development increments',delta),
      ('S18','Critical validation integrity tests',pd.read_csv(ROOT/'04_validation/critical_integrity_tests.csv')),
      ('S19','Linear, age-spline and limited nonlinear model comparisons',perf[perf.model.str.contains('spline|RSF')]),
      ('S20','Cross-validated monthly-grid IBS and actual follow-up support',ibs_summary)
    ]
    settings=json.loads((ROOT/'02_extraction/extraction_effective_settings.json').read_text())
    for name,title,frame in tables:
        clean=frame.copy()
        for col in clean.select_dtypes(include=['object']).columns:
            clean[col]=clean[col].map(lambda v:str(v).replace(str(ROOT.parent),'.') if isinstance(v,str) else v)
        s=add_sheet(sw,name,f'Supplementary Table {name}. {title}.',clean)
        if name=='S12':
            for col,width in {'A':48,'B':38,'C':52,'D':26,'E':26,'F':24,'G':12}.items():s.column_dimensions[col].width=width
            s.append([]); s.append(['Configuration domain','Item','Final reported setting'])
            for domain,items in [('Extraction settings',settings['settings']),('Software',settings['distribution']),('Other audited settings',{k:v for k,v in settings.items() if k not in ['settings','distribution','featureClass']})]:
                for key,v in items.items():
                    s.append([domain,key,value(v)]); s.cell(s.max_row,3).alignment=Alignment(wrap_text=True); s.row_dimensions[s.max_row].height=min(409,max(30,int(np.ceil(len(str(value(v)))/50))*13+8))
    sw.save(OUT/'Supplementary_Tables.xlsx')
    (OUT/'TABLE_AND_FIGURE_SOURCE_INDEX.json').write_text(json.dumps({'main_tables':['Table 1','Table 2','Table 3','Table 4'],'supplementary_tables':[{ 'number':n,'title':title} for n,title,_ in tables]},indent=2),encoding='utf-8')
    print('Publication tables and calibration source data saved.',flush=True)

if __name__=='__main__':main()
